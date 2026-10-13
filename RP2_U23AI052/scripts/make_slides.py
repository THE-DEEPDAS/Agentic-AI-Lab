#!/usr/bin/env python3
"""Build the 10-slide deck (python-pptx, native charts).

    python scripts/make_slides.py --run-dir results/<real run> --out slides/RP2_slides.pptx
    python scripts/make_slides.py                                   # no run yet: results slides say 'pending'

Result slides are filled from <run-dir>/summary.json. With the mock run they carry a red
'MOCK MODEL - NOT A RESULT' banner.
"""
import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from rp2.aggregate import LABELS  # noqa: E402
from rp2.common import ROOT, load_config, load_pool  # noqa: E402
from rp2.report_text import describe, failure_share  # noqa: E402
from rp2.scoring import FAILURES  # noqa: E402

NAVY = RGBColor(0x1F, 0x2A, 0x6B)
INK = RGBColor(0x1B, 0x1B, 0x1F)
MUTED = RGBColor(0x5F, 0x63, 0x73)
LIGHT = RGBColor(0xF0, 0xF1, 0xF5)
RED = RGBColor(0xB3, 0x26, 0x1E)
BLUE = RGBColor(0x2E, 0x6F, 0xD8)
ORANGE = RGBColor(0xE0, 0x7B, 0x1A)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
PALETTE = [BLUE, ORANGE, GREEN, RGBColor(0x9B, 0x3F, 0xB5), RGBColor(0xC2, 0x3B, 0x5A), RGBColor(0x6B, 0x7A, 0x8F),
           RGBColor(0x8C, 0x6D, 0x1F), RGBColor(0x1F, 0x8F, 0x9B), RGBColor(0x55, 0x55, 0x55)]
FONT = "Calibri"


class Deck:
    def __init__(self, mock):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.mock = mock
        self.n = 0

    def slide(self, title, notes="", result_slide=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        self.text(s, title, 0.6, 0.35, 12.1, 0.9, size=30, bold=True, color=NAVY)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.2), Inches(1.4), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        self.text(s, f"AI465 Agentic AI  |  RP2  |  U23AI052   {self.n}", 0.6, 7.05, 12.1, 0.3, size=10, color=MUTED)
        if self.mock and result_slide:
            self.text(s, "MOCK MODEL - NOT A RESULT", 8.2, 0.4, 4.6, 0.5, size=18, bold=True, color=RED, align=PP_ALIGN.RIGHT)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def text(self, s, text, x, y, w, h, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             italic=False):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        lines = text if isinstance(text, list) else [text]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ln
            r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, FONT
            r.font.color.rgb = color
        return tb

    def bullets(self, s, items, x, y, w, h, size=18, gap=8):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            head, _, rest = it.partition("||")
            r = p.add_run()
            r.text = "•  " + head
            r.font.size, r.font.name, r.font.bold = Pt(size), FONT, bool(rest)
            r.font.color.rgb = INK
            if rest:
                r2 = p.add_run()
                r2.text = rest
                r2.font.size, r2.font.name, r2.font.bold = Pt(size), FONT, False
                r2.font.color.rgb = INK
        return tb

    def box(self, s, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        b.fill.solid()
        b.fill.fore_color.rgb = fill
        if line:
            b.line.color.rgb = line
        else:
            b.line.fill.background()
        b.shadow.inherit = False
        return b

    def table(self, s, rows, x, y, w, col_w, size=12, row_h=0.36):
        shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
        t = shp.table
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = t.cell(i, j)
                c.text = str(v)
                p = c.text_frame.paragraphs[0]
                p.runs[0].font.size = Pt(size)
                p.runs[0].font.name = FONT
                p.runs[0].font.bold = (i == 0)
                p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
                c.fill.solid()
                c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
                c.margin_left = c.margin_right = Inches(0.06)
                c.margin_top = c.margin_bottom = Inches(0.02)
        return t

    def chart(self, s, kind, cats, series, x, y, w, h, legend=True, pct=True, title=None, stacked_colors=None):
        cd = CategoryChartData()
        cd.categories = cats
        for name, vals in series:
            cd.add_series(name, vals)
        gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.font.size, ch.font.name = Pt(12), FONT
        ch.has_legend = legend
        if legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        if title:
            ch.has_title = True
            ch.chart_title.text_frame.text = title
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        else:
            ch.has_title = False
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xE3)
        va.format.line.fill.background()
        if pct:
            va.maximum_scale, va.minimum_scale = 100, 0
            va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0"%"', False
        plot = ch.plots[0]
        plot.gap_width = 70
        for i, ser in enumerate(plot.series):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = (stacked_colors or PALETTE)[i % len(PALETTE)]
        return ch


def pending(deck, s, what):
    deck.box(s, 0.9, 2.0, 11.5, 3.6, fill=LIGHT)
    deck.text(s, "Result pending", 1.2, 2.3, 10.9, 0.7, size=28, bold=True, color=NAVY)
    deck.text(s, [what, "", "Run: python scripts/run_experiment.py --tiers 1 2 3",
                  "then: python scripts/aggregate.py --run-dir results/<run>",
                  "then: python scripts/make_slides.py --run-dir results/<run>"], 1.2, 3.1, 10.9, 2.4, size=18, color=INK)


def build(run_dir, out):
    cfg = load_config()
    pool = load_pool(cfg)
    summary = fcase = None
    if run_dir and os.path.exists(os.path.join(run_dir, "summary.json")):
        summary = json.load(open(os.path.join(run_dir, "summary.json"), encoding="utf-8"))
        fp = os.path.join(run_dir, "failure_case.json")
        fcase = json.load(open(fp, encoding="utf-8")) if os.path.exists(fp) else None
    mock = bool(summary and summary["mock_run"])
    S = summary["conditions"] if summary else {}
    D = Deck(mock)

    # 1 title
    s = D.slide("", notes="(0:00-0:30) Introduce the question: does a longer tool list make an LLM worse at choosing the right MCP tool, and what does it cost? "
                          "Individual project, roll number U23AI052.")
    D.text(s, "Does a longer tool list hurt MCP tool selection?", 0.8, 2.2, 11.7, 1.4, size=44, bold=True, color=NAVY)
    D.text(s, "5 vs 20 vs 50 tools in context on a 200-task MCPToolBench++ subset", 0.8, 3.7, 11.7, 0.6, size=24, color=INK)
    D.text(s, "AI465 Agentic AI  |  Research Project, Phase 2: implementation and evaluation", 0.8, 4.6, 11.7, 0.5, size=18, color=MUTED)
    D.text(s, "U23AI052  |  Individual work  |  SVNIT Surat", 0.8, 5.1, 11.7, 0.5, size=18, color=MUTED)
    if mock:
        D.text(s, "MOCK MODEL - NOT A RESULT", 0.8, 6.2, 11.7, 0.6, size=22, bold=True, color=RED)

    # 2 objective and deviations
    s = D.slide("Objective, and what changed from Phase 1",
                notes="(0:30-1:30) Read the objective, then be upfront about the deviations. The public release has only single-step data for these three domains, "
                      "so multi-step and AST DAG are dropped. The GPT-4o judge is replaced by a deterministic scorer that I checked against my own labels.")
    D.box(s, 0.6, 1.5, 4.6, 5.3)
    D.text(s, "Approved objective", 0.8, 1.6, 4.2, 0.5, size=20, bold=True, color=NAVY)
    D.text(s, ["Effect of tool-count context (5 vs 20 vs 50 tools) on MCP tool-selection accuracy and failure composition, on a stratified 200-task "
               "MCPToolBench++ subset (search, map, finance).", "",
               "Metrics: AST accuracy, tokens, calls and latency per task, plus a failure taxonomy.", "",
               "Fills the paper's gaps: no cost, no variance, no failure analysis."], 0.8, 2.2, 4.2, 4.4, size=16)
    D.text(s, "Declared deviations", 5.5, 1.5, 7.2, 0.5, size=20, bold=True, color=NAVY)
    D.bullets(s, [
        "D1 Single-step only.|| Public data has no multi-step items for these domains, so no AST DAG.",
        "D2 No Pass@K.|| Needs live MCP servers and credentials.",
        "D3 Deterministic scorer|| instead of the GPT-4o judge; audited against my labels.",
        "D4 Distractors|| from an 87-tool pool; native lists are only 1, 5 or 32 tools.",
        "D5 Matched budget|| = identical prompts + 5-tool baseline with k=8 votes.",
        "D6 Missing required parameter|| counts as an error (stricter).",
    ], 5.5, 2.1, 7.3, 4.7, size=16, gap=9)

    # 3 design
    s = D.slide("Design: the tool list is the only thing that changes",
                notes="(1:30-2:30) Nested design: the 5-tool context is a subset of the 20-tool one, which is a subset of the 50-tool one. "
                      "Same prompt files, same model, same decoding. Three ablations isolate one factor each: position, length, budget.")
    for i, (n, wd, col) in enumerate([(50, 9.6, RGBColor(0xC9, 0xD6, 0xF2)), (20, 6.2, RGBColor(0x8F, 0xAE, 0xE8)), (5, 2.6, BLUE)]):
        D.box(s, 0.7, 1.55 + i * 0.52, wd, 0.44, fill=col, shape=MSO_SHAPE.RECTANGLE)
        D.text(s, f"{n} tools", 0.8, 1.57 + i * 0.52, 2.2, 0.4, size=14, bold=True, color=INK if n != 5 else RGBColor(255, 255, 255))
    D.text(s, "gold tool + up to 2 labelled similar tools are always present; distractors come from one seeded permutation per task",
           0.7, 3.15, 11.9, 0.4, size=13, color=MUTED, italic=True)
    D.table(s, [["Condition", "What differs", "Isolates"],
                ["c05 / c20 / c50", "5, 20, 50 tools in context", "effect of tool count (the objective)"],
                ["c50_goldfirst / goldlast", "gold tool placed first or last at 50 tools", "position of the right tool"],
                ["c05_padded", "5 tools + neutral text to the 50-tool length", "context length vs number of choices"],
                ["c05_sc", "5 tools, 8 samples, majority vote", "matched token budget"]],
            0.7, 3.8, 11.9, [3.0, 4.9, 4.0], size=14, row_h=0.5)
    D.text(s, "Same system and user prompt files in every condition (a test asserts it). Every run: 200 tasks x 3 repetitions, temperature and seed pinned.",
           0.7, 6.4, 11.9, 0.5, size=14, color=INK)

    # 4 implementation
    s = D.slide("What I built",
                notes="(2:30-3:30) Walk the pipeline left to right. Everything is mine except the benchmark's task data and labels. "
                      "A failed API call becomes an api_error record and is never a crash; interrupted runs resume. unit tests including an oracle test.")
    steps = [("Subset builder", "200 tasks, 87-tool pool"), ("Condition builder", "nested tool lists"), ("Runner", "model adapters, retries, resume"),
             ("Traces", "prompt, tools, response, usage"), ("Scorer", "tool_ok, param_ok, 9 failure labels"),
             ("Aggregate", "mean±SD, bootstrap, figures")]
    for i, (h, sub) in enumerate(steps):
        x = 0.6 + i * 2.1
        D.box(s, x, 1.7, 1.85, 1.6, fill=LIGHT)
        D.text(s, h, x + 0.05, 1.8, 1.75, 0.5, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        D.text(s, sub, x + 0.05, 2.3, 1.75, 0.9, size=12, color=INK, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            D.text(s, "→", x + 1.83, 2.2, 0.3, 0.5, size=22, bold=True, color=NAVY)
    D.bullets(s, [
        "Own code:|| runner, condition builder, scorer, statistics, report and slide generators, tests.",
        "Borrowed:|| the benchmark's task data and labels (commit 2bd35e5); its judge rules re-implemented, no code imported.",
        "Robustness:|| retries with backoff, rate limit, resume, jobs interleaved across conditions so model drift hits all of them.",
        "Checked:|| oracle test (200 gold labels score 100%), perturbation tests, 20-task audit against my own labels.",
    ], 0.7, 3.7, 12.0, 3.2, size=17, gap=10)

    # 5 protocol
    s = D.slide("Experimental protocol",
                notes="(3:30-4:15) These are the rubric requirements. Each one maps to a file in the repository. Planned token cost is in the README.")
    D.table(s, [["Requirement", "How it is met"],
                ["Repetition", "3 full passes per condition; mean ± SD over runs; paired bootstrap over tasks"],
                ["Pinned configuration", "configs/config.yaml, manifest.json per run: model string asked and returned, temperature, top-p, seed, dates"],
                ["Prompts as artefacts", "prompts/*.txt; conditions differ only in the tool list (padding arm adds neutral text)"],
                ["Metrics", "tool and AST accuracy; model calls, input and output tokens, latency per task"],
                ["Traces", "every call: full prompt, tool list, raw response, usage, score, failure label"],
                ["Judging", "no LLM judge; deterministic scorer checked on 20 tasks against my labels"],
                ["Reproducibility", "one command per table; README states runtime and token cost"]],
            0.7, 1.6, 11.9, [3.0, 8.9], size=14, row_h=0.6)

    # 6 results: accuracy
    s = D.slide("Results: accuracy vs number of tools",
                notes="(4:15-5:30) Read the three AST numbers, then the paired bootstrap interval for 5 to 50 tools. Say whether it is outside sampling noise. "
                      "Then point to the tool-selection series: is the loss selection or argument filling?", result_slide=True)
    if summary and all(c in S for c in ("c05", "c20", "c50")):
        core = ["c05", "c20", "c50"]
        D.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, [LABELS[c] for c in core],
                [("Tool selection %", [round(S[c]["tool_acc"]["mean"] * 100, 1) for c in core]),
                 ("AST (tool + parameters) %", [round(S[c]["ast_acc"]["mean"] * 100, 1) for c in core])],
                0.6, 1.5, 7.4, 5.3)
        c = (summary["comparisons"].get("c05->c50") or {}).get("ast_ok")
        sd = [S[x]["ast_acc"]["sd"] * 100 for x in core]
        lines = [f"AST accuracy", f"{S['c05']['ast_acc']['mean'] * 100:.1f}%  ->  {S['c20']['ast_acc']['mean'] * 100:.1f}%  ->  {S['c50']['ast_acc']['mean'] * 100:.1f}%",
                 f"run-to-run SD {sd[0]:.1f} / {sd[1]:.1f} / {sd[2]:.1f} pp"]
        D.box(s, 8.3, 1.6, 4.5, 2.1)
        D.text(s, lines, 8.45, 1.7, 4.2, 1.9, size=18, bold=False)
        if c:
            D.box(s, 8.3, 3.95, 4.5, 2.1)
            D.text(s, ["5 to 50 tools, paired bootstrap", f"{c['mean_diff'] * 100:+.1f} pp", f"95% CI [{c['ci95'][0] * 100:+.1f}, {c['ci95'][1] * 100:+.1f}]",
                       "outside noise" if c["excludes_zero"] else "within noise"], 8.45, 4.05, 4.2, 1.9, size=18)
        D.text(s, f"Mean over {S['c05']['n_runs']} runs x {S['c05']['n_tasks']} tasks.", 8.3, 6.3, 4.5, 0.4, size=12, color=MUTED)
    else:
        pending(D, s, "Accuracy at 5, 20 and 50 tools with run-to-run SD and the paired bootstrap interval.")

    # 7 results: cost
    s = D.slide("Results: what the extra tools cost",
                notes="(5:30-6:15) Tokens per task grow with the tool list; the question is whether accuracy bought with tokens could be had more cheaply. "
                      "The self-consistency arm gives the 5-tool baseline about the same token budget as 50 tools.", result_slide=True)
    if summary and all(c in S for c in ("c05", "c20", "c50")):
        cs = [c for c in ("c05", "c20", "c50", "c05_sc") if c in S]
        D.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, [LABELS[c] for c in cs],
                [("Input tokens per task", [round(S[c]["input_tokens"]["mean"]) for c in cs])], 0.6, 1.5, 6.2, 5.3,
                legend=False, pct=False, title="Input tokens per task")
        rows = [["Condition", "AST %", "Latency s", "Tok / correct"]]
        for c in cs:
            rows.append([LABELS[c], f"{S[c]['ast_acc']['mean'] * 100:.1f}", f"{S[c]['latency_s']['mean']:.2f}",
                         "n/a" if not S[c]["tokens_per_correct"] else f"{S[c]['tokens_per_correct']:.0f}"])
        D.table(s, rows, 7.1, 1.7, 5.7, [2.3, 0.9, 1.2, 1.3], size=13, row_h=0.5)
        D.text(s, "Calls per task: 1 in every condition except self-consistency (k votes).", 7.1, 4.6, 5.7, 0.8, size=14, color=MUTED)
    else:
        pending(D, s, "Tokens, calls and latency per task, and accuracy against budget.")

    # 8 ablations
    s = D.slide("Ablations: what explains the effect?",
                notes="(6:15-7:00) Position arm: is it where the right tool sits? Padding arm: is it context length rather than number of choices? "
                      "Each row is a paired difference with a bootstrap interval.", result_slide=True)
    comps = summary["comparisons"] if summary else {}
    rowsdef = [("c50->c50_goldfirst", "50 tools: gold first vs shuffled", "position"), ("c50->c50_goldlast", "50 tools: gold last vs shuffled", "position"),
               ("c05->c05_padded", "5 tools padded to 50-tool length vs plain 5", "context length"),
               ("c05->c05_sc", "5 tools, 8 votes vs plain 5", "budget"), ("c05_sc->c50", "50 tools vs 5 tools with 8 votes", "matched budget")]
    rows = [["Comparison (AST accuracy)", "Factor", "Diff (pp)", "95% CI (pp)", "Beyond noise"]]
    have = False
    for k, lab, fac in rowsdef:
        c = (comps.get(k) or {}).get("ast_ok")
        if c:
            have = True
            rows.append([lab, fac, f"{c['mean_diff'] * 100:+.1f}", f"[{c['ci95'][0] * 100:+.1f}, {c['ci95'][1] * 100:+.1f}]", "yes" if c["excludes_zero"] else "no"])
    if have:
        D.table(s, rows, 0.7, 1.7, 11.9, [5.2, 1.9, 1.4, 2.2, 1.2], size=14, row_h=0.6)
        D.text(s, "Differences are computed per task (mean over runs) and bootstrapped over the 200 tasks.", 0.7, 6.2, 11.9, 0.5, size=14, color=MUTED)
    else:
        pending(D, s, "Position, padding and matched-budget comparisons with bootstrap intervals (tiers 2 and 3).")

    # 9 failures
    s = D.slide("Why it fails: taxonomy and a traced case",
                notes="(7:00-7:45) Failure labels are exclusive and computed from the trace. Then read one failure from its trace: query, gold tool, the tool it picked, "
                      "and the tool descriptions. State the cause in your own words.", result_slide=True)
    F = summary["failures"] if summary else {}
    if summary and "c05" in F and "c50" in F:
        cats = [c for c in FAILURES if F["c05"]["counts"][c] + F["c50"]["counts"][c] > 0]
        series = [(c.replace("_", " "), [round(F[x]["counts"][c] / F[x]["n"] * 100, 1) for x in ("c05", "c50")]) for c in cats]
        D.chart(s, XL_CHART_TYPE.COLUMN_STACKED, ["5 tools", "50 tools"], series, 0.6, 1.5, 5.6, 5.3, pct=False,
                title="Share of all attempts (%)")
        if fcase:
            b = fcase["c50_failure"]
            sb = b["score"]
            picked = sb["pred_name"]
            D.box(s, 6.5, 1.6, 6.3, 5.2)
            D.text(s, [f"Task {b['uuid'][:8]} ({b['domain']}): solved in every 5-tool run, failed in every 50-tool run",
                       "", f"Query: {b['user'][:150]}", f"Gold: {b['label']['name']}", f"Called: {picked}  ->  {sb['failure'].replace('_', ' ')}",
                       f"Gold at position {b['gold_position']} of {b['n_tools']}", "",
                       f"Picked tool: {describe(pool, picked, 120) if picked else '-'}", f"Gold tool: {describe(pool, b['label']['name'], 120)}",
                       "", "Cause: (state after reading the trace)"], 6.65, 1.7, 6.0, 5.0, size=14)
        else:
            D.text(s, "No task was solved in every 5-tool run and failed in every 50-tool run.", 6.5, 3.0, 6.3, 1.0, size=16)
    else:
        pending(D, s, "Failure composition at 5 and 50 tools, plus one failure case read from its trace.")

    # 10 limits and demo
    s = D.slide("Limitations, next steps, and the demo",
                notes="(7:45-8:00 then LIVE DEMO) Run: python scripts/demo.py --domain map   (or --task-index N). It prints the query, gold tool, then the same task at 5, 20 and 50 "
                      "tools: tool called, pass/fail with failure label, tokens, latency. Traces are opened from results/demo/traces. "
                      "If the API is slow, play the recorded run.")
    D.text(s, "Limitations", 0.6, 1.5, 6.0, 0.5, size=20, bold=True, color=NAVY)
    D.bullets(s, ["One model, one temperature, single-step tasks.",
                  "Labels are the benchmark's; a valid alternative tool can be scored wrong.",
                  "Distractors are mostly cross-domain, so easier than a real marketplace.",
                  "Lenient parameter matching, as in the benchmark.",
                  "No Pass@K or live MCP execution."], 0.6, 2.1, 6.0, 4.5, size=16, gap=8)
    D.text(s, "With four more weeks", 6.9, 1.5, 6.0, 0.5, size=20, bold=True, color=NAVY)
    D.bullets(s, ["Several models and sizes.",
                  "Distractors chosen by similarity to the gold tool.",
                  "A retrieval baseline that shows only the top-k tools.",
                  "Multi-step chains built from labelled calls.",
                  "Live MCP servers for Pass@K."], 6.9, 2.1, 6.0, 3.0, size=16, gap=8)
    D.box(s, 6.9, 5.2, 5.9, 1.5, fill=NAVY)
    D.text(s, ["Live demo", "python scripts/demo.py --domain map"], 7.05, 5.3, 5.6, 1.3, size=20, bold=True, color=RGBColor(255, 255, 255))

    os.makedirs(os.path.dirname(out), exist_ok=True)
    D.prs.save(out)
    return D.n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run-dir", default=None)
    ap.add_argument("--out", default=os.path.join(ROOT, "slides", "RP2_slides.pptx"))
    a = ap.parse_args()
    n = build(a.run_dir, a.out)
    print(f"wrote {a.out}: {n} slides")
    return 0


if __name__ == "__main__":
    sys.exit(main())
